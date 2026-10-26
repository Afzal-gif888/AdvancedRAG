"""Turn raw bytes into page-aware text plus an extraction report. Pure functions, run in a worker thread.

``parse_document`` is the full interface (pages + a report of how every page was obtained and what could
not be processed). ``parse`` keeps the original ``bytes -> List[Page]`` contract for existing callers.

Sources may be ``bytes``, a filesystem path, or a binary file object, so a large upload can be parsed from
the spooled file on disk instead of being held in memory several times.
"""
import csv
import io
import logging
import os
import re
import time
import zipfile
from dataclasses import dataclass, field
from typing import Any, BinaryIO, Dict, List, Optional, Tuple, Union, cast

from advancerag_common.observability import METRICS
from rag_core import errors, ocr, pdftables, textclean

log = logging.getLogger("advancerag.rag.parsing")

TEXT_TYPES = {"text/plain", "text/markdown", "text/csv"}
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CSV = "text/csv"
MARKDOWN = "text/markdown"
IMAGE_TYPES = {"image/png", "image/jpeg"}

Source = Union[bytes, bytearray, str, "os.PathLike[str]", BinaryIO]


@dataclass
class Page:
    number: Optional[int]  # 1-based for PDFs; None for formats without pages
    text: str
    method: str = "native"  # native | ocr
    confidence: Optional[float] = None  # engine-reported mean, OCR pages only
    # (char offset in `text`, level 1..6, title): only where the format exposes reliable structure
    headings: List[Tuple[int, int, str]] = field(default_factory=list)


@dataclass
class ExtractionReport:
    """How the text was obtained, and everything that could not be."""

    source_type: str
    total_pages: Optional[int] = None  # None for formats without pages
    native_pages: int = 0
    ocr_pages: int = 0
    blank_pages: List[int] = field(default_factory=list)  # nothing to extract (and nothing was lost)
    unprocessed_pages: List[int] = field(default_factory=list)  # image-bearing / skipped: content NOT searchable
    failed_pages: List[int] = field(default_factory=list)
    warnings: List[Dict[str, Any]] = field(default_factory=list)
    ocr_engine: Optional[str] = None
    ocr_confidence_mean: Optional[float] = None
    cleanup: Dict[str, int] = field(default_factory=dict)
    # Tables turned into structured rows: {page, kind (lines|text|ocr_layout), rows, columns, header, continued}
    tables: List[Dict[str, Any]] = field(default_factory=list)

    @property
    def partial(self) -> bool:
        """True when some content could not be made searchable (the document is still usable)."""
        return bool(self.unprocessed_pages or self.failed_pages)

    @property
    def method(self) -> str:
        if self.ocr_pages and self.native_pages:
            return "mixed"
        return "ocr" if self.ocr_pages else "native"

    def warn(self, page: Optional[int], code: str, message: str) -> None:
        self.warnings.append({"page": page, "code": code, "message": message})

    def to_dict(self) -> Dict[str, Any]:
        return {
            "source_type": self.source_type,
            "method": self.method,
            "partial": self.partial,
            "total_pages": self.total_pages,
            "native_pages": self.native_pages,
            "ocr_pages": self.ocr_pages,
            "blank_pages": self.blank_pages,
            "unprocessed_pages": self.unprocessed_pages,
            "failed_pages": self.failed_pages,
            "ocr_engine": self.ocr_engine,
            "ocr_confidence_mean": None if self.ocr_confidence_mean is None else round(self.ocr_confidence_mean, 3),
            "cleanup": self.cleanup,
            "tables": self.tables[:50],
            "tables_truncated": len(self.tables) > 50,
            "warnings": self.warnings[:50],
            "warnings_truncated": len(self.warnings) > 50,
        }


@dataclass
class ParsedDocument:
    pages: List[Page]
    report: ExtractionReport


# ---------------------------------------------------------------- text normalisation

_WS = re.compile(r"[ \t\f\v]+")
_BLANKS = re.compile(r"\n{3,}")


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(_WS.sub(" ", line).strip() for line in text.split("\n"))
    return _BLANKS.sub("\n\n", text).strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


# ---------------------------------------------------------------- source helpers

def _read_bytes(source: Source) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, os.PathLike)):
        with open(source, "rb") as f:
            return f.read()
    source.seek(0)
    return source.read()


def _stream(source: Source) -> BinaryIO:
    """A seekable binary stream over any source (paths are opened; bytes wrapped)."""
    if isinstance(source, (bytes, bytearray)):
        return io.BytesIO(source)
    if isinstance(source, (str, os.PathLike)):
        return open(source, "rb")
    source.seek(0)
    return source


def _pdfium_source(source: Source) -> Any:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, os.PathLike)):
        return str(source)
    source.seek(0)
    return source


# ---------------------------------------------------------------- tables (DOCX, CSV)

def _row_text(header: Optional[List[str]], cells: List[str]) -> str:
    """One self-describing line per table/CSV row: `Battery life | 14 months` becomes
    `Property: Battery life | Value: 14 months` when a header exists, so a chunk taken from the
    middle of a long table still says what each value means."""
    if not header:
        return " | ".join(c for c in cells if c)
    parts = []
    for i, value in enumerate(cells):
        if value:
            name = header[i] if i < len(header) and header[i] else f"column {i + 1}"
            parts.append(f"{name}: {value}")
    return " | ".join(parts)


_NUMBER = re.compile(r"^[-+]?[\d.,\s%€$£]+$")


def _usable_header(cells: List[str]) -> Optional[List[str]]:
    """The first row is used as column names only if it looks like one: >= 2 distinct,
    non-empty names that are not all numbers (an all-numeric first row is data)."""
    names = [c.strip() for c in cells]
    if len(names) < 2 or any(not n for n in names) or len(set(names)) != len(names):
        return None
    if all(_NUMBER.match(n) for n in names):
        return None
    return names


def _drop_repeated_headers(header: List[str], body: List[List[str]]) -> List[List[str]]:
    """Tables exported page by page repeat their header row; such a row is not data."""
    key = [c.strip().lower() for c in header]
    return [r for r in body if [c.strip().lower() for c in r[: len(key)]] != key or len(r) > len(key)]


def _docx_table(table) -> List[str]:
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:  # horizontally merged cells repeat: keep each underlying cell once
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(" ".join(cell.text.split()))
        rows.append(cells)
    if not rows:
        return []
    header = _usable_header(rows[0]) if len(rows) > 1 else None
    body = rows[1:] if header else rows
    if header:
        body = _drop_repeated_headers(header, body)
    lines = [_row_text(header, cells) for cells in body]
    return ([" | ".join(header)] if header else []) + [line for line in lines if line]


def _check_zip_safety(source: Source, cfg) -> None:
    """Refuse a DOCX package that expands absurdly (zip bomb) before anything decompresses it."""
    stream = _stream(source)
    close = isinstance(source, (str, os.PathLike))
    try:
        stream.seek(0)
        with zipfile.ZipFile(stream) as zf:
            infos = zf.infolist()
            if len(infos) > 20_000:
                raise errors.ServiceError(errors.PARSE_FAILED, "The document structure is not safe to process.", 422)
            total = sum(i.file_size for i in infos)
            packed = max(sum(i.compress_size for i in infos), 1)
            if total > cfg.MAX_ZIP_UNCOMPRESSED_BYTES or total / packed > cfg.MAX_ZIP_RATIO:
                raise errors.ServiceError(errors.PARSE_FAILED, "The document is too large when unpacked.", 422)
    except zipfile.BadZipFile as exc:
        raise errors.parse_failed() from exc
    finally:
        if close:
            stream.close()
        else:
            stream.seek(0)


def _docx_text(source: Source) -> Tuple[str, List[Tuple[int, str]]]:
    """Paragraphs and tables in document order (a table stays under its heading).
    Returns (text, [(level, title)] of heading paragraphs in order)."""
    import docx
    from docx.table import Table

    stream = _stream(source)
    try:
        document = docx.Document(stream)
        parts: List[str] = []
        headings: List[Tuple[int, str]] = []
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                parts.extend(_docx_table(block))
            else:
                parts.append(block.text)
                level = _heading_level(getattr(block.style, "name", "") or "")
                if level and block.text.strip():
                    headings.append((level, " ".join(block.text.split())))
        return "\n".join(parts), headings
    finally:
        if isinstance(source, (str, os.PathLike)):
            stream.close()


def _heading_level(style_name: str) -> int:
    m = re.match(r"^Heading\s*(\d)$", style_name.strip(), re.I)
    if m:
        return max(1, min(int(m.group(1)), 6))
    return 1 if style_name.strip().lower() == "title" else 0


def locate_headings(text: str, wanted: List[Tuple[int, str]]) -> List[Tuple[int, int, str]]:
    """Find each (level, title) as a whole line of ``text``, in order. Titles that cannot be found are
    skipped: a heading is never invented."""
    found, pos = [], 0
    for level, title in wanted:
        m = re.compile(rf"(?m)^{re.escape(title)}[ \t]*$").search(text, pos)
        if m:
            found.append((m.start(), level, title))
            pos = m.end()
    return found


_MD_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*#*[ \t]*$")


def markdown_headings(text: str) -> List[Tuple[int, int, str]]:
    """ATX headings (``# Title``) outside fenced code blocks."""
    out, offset, in_fence = [], 0, False
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            in_fence = not in_fence
        elif not in_fence:
            m = _MD_HEADING.match(line)
            if m:
                out.append((offset, len(m.group(1)), m.group(2).strip()))
        offset += len(line) + 1
    return out


# ---------------------------------------------------------------- CSV

def _csv_rows(text: str) -> Optional[List[List[str]]]:
    """Non-empty rows of a delimited table, or None when the text does not look like one."""
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error:
        return None
    rows = [r for r in rows if any(r)]
    if len(rows) < 2 or max(len(r) for r in rows) < 2:
        return None
    return rows


def _csv_text(text: str) -> str:
    """Rows rendered as `column: value | ...` lines. Falls back to the raw text when the file
    does not look like a delimited table."""
    rows = _csv_rows(text)
    if rows is None:
        return text
    header = _usable_header(rows[0])
    body = rows[1:] if header else rows
    if header:
        body = _drop_repeated_headers(header, body)
    return "\n".join(line for line in (_row_text(header, r) for r in body) if line)


@dataclass
class Table:
    columns: List[str]
    rows: List[List[str]]


def parse_table(data: Source, max_rows: int, max_columns: int) -> Optional[Table]:
    """The complete CSV as a table (for exact, whole-dataset analysis), with the same dialect and
    header detection as the text rendering. None when it is not a table or exceeds the limits."""
    rows = _csv_rows(_decode(_read_bytes(data)))
    if rows is None:
        return None
    header = _usable_header(rows[0])
    body = rows[1:] if header else rows
    if header:
        body = _drop_repeated_headers(header, body)
    width = max(len(header or []), max((len(r) for r in body), default=0))
    if not body or len(body) > max_rows or width > max_columns:
        return None
    columns = [(header[i] if header and i < len(header) else f"column {i + 1}") for i in range(width)]
    return Table(columns, [(r + [""] * (width - len(r)))[:width] for r in body])


# ---------------------------------------------------------------- PDF

def _page_has_images(page) -> bool:
    """Does the page draw at least one image XObject? (cheap check; inline images are not detected)"""
    try:
        resources = page.get("/Resources")
        xobjects = resources.get_object().get("/XObject") if resources is not None else None
        if xobjects is None:
            return False
        for obj in xobjects.get_object().values():
            if obj.get_object().get("/Subtype") == "/Image":
                return True
    except Exception:  # malformed resource trees: treat as "unknown", i.e. no images
        return False
    return False


def _looks_blank(img) -> bool:
    lo, hi = img.convert("L").getextrema()
    return (hi - lo) < 12


def _parse_pdf(source: Source, cfg, report: ExtractionReport) -> List[Page]:
    from pypdf import PdfReader

    stream = _stream(source)
    try:
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise errors.ServiceError(errors.PARSE_FAILED, "Password-protected PDFs are not supported.", 422)
        total = len(reader.pages)
        if total > cfg.MAX_PDF_PAGES:
            raise errors.ServiceError(errors.DOCUMENT_TOO_LARGE, f"The PDF has {total} pages; the limit is {cfg.MAX_PDF_PAGES}.", 413)
        report.total_pages = total

        texts: List[Optional[str]] = []
        images: List[bool] = []
        parse_deadline = time.monotonic() + cfg.PDF_PARSE_TIMEOUT_SECONDS
        for i, p in enumerate(reader.pages):
            if time.monotonic() > parse_deadline:
                # Cooperative budget, checked between pages: a thread cannot be killed, so one page that never returns is
                # not interruptible here (that needs the parser in a separate process; see docs/verification/SECURITY_AUDIT.md).
                texts.append(None)
                images.append(False)
                report.failed_pages.append(i + 1)
                report.warn(i + 1, "PARSE_TIME_BUDGET", f"Reading this PDF took longer than {cfg.PDF_PARSE_TIMEOUT_SECONDS:.0f} s; this page was not read.")
                continue
            try:
                raw = p.extract_text() or ""
                texts.append(normalize(textclean.clean_chars(raw)))
            except Exception:
                texts.append(None)  # this page failed natively; the others are still fine
                report.failed_pages.append(i + 1)
                report.warn(i + 1, "NATIVE_EXTRACT_FAILED", "Text could not be extracted from this page.")
            images.append(_page_has_images(p))
        if total and len(report.failed_pages) == total:
            raise errors.parse_failed()

        methods = ["native"] * total
        confs: List[Optional[float]] = [None] * total
        candidates = [i for i, t in enumerate(texts) if t is not None and (not t or (len(t) < cfg.OCR_MIN_NATIVE_CHARS and images[i]))]
        if candidates and cfg.OCR_ENABLED:
            _ocr_pages(source, cfg, report, texts, methods, confs, images, candidates)
        else:
            pdf = None
            for i in candidates:
                if not texts[i] and images[i]:
                    # An image page may be a blank scan: look at a thumbnail before declaring content lost.
                    try:
                        if pdf is None:
                            pdf = ocr.open_pdfium(_pdfium_source(source))
                        if _looks_blank(ocr.render_pdf_page(pdf, i, 36, 400_000)):
                            continue  # counted as a blank page below
                    except Exception:
                        pass
                    report.unprocessed_pages.append(i + 1)
                    report.warn(i + 1, "NO_TEXT_OCR_DISABLED", "This page contains only an image; OCR is not enabled, so it is not searchable.")
            if pdf is not None:
                try:
                    pdf.close()
                except Exception:
                    pass

        if cfg.PDF_TABLES != "off":
            _structure_tables(source, cfg, report, texts, methods)

        # hygiene across the whole document (needs all pages: repeated headers/footers)
        cleanup = {"hyphen_joins": 0, "margin_lines_removed": 0}
        if cfg.TEXT_CLEANUP:
            idx = [i for i, t in enumerate(texts) if t]
            cleaned, removed = textclean.remove_repeated_margins([cast(str, texts[i]) for i in idx])
            cleanup["margin_lines_removed"] = removed
            for i, t in zip(idx, cleaned):
                texts[i] = t
            for i in idx:
                t, joins = textclean.dehyphenate(cast(str, texts[i]))
                cleanup["hyphen_joins"] += joins
                texts[i] = t
            if cfg.UNWRAP_LINES:
                for i in idx:
                    texts[i] = textclean.unwrap_lines(cast(str, texts[i]))
        report.cleanup = cleanup

        pages: List[Page] = []
        for i, t in enumerate(texts):
            number = i + 1
            if t:
                pages.append(Page(number, normalize(t), methods[i], confs[i]))
                if methods[i] == "ocr":
                    report.ocr_pages += 1
                else:
                    report.native_pages += 1
            elif t == "" and number not in report.unprocessed_pages and number not in report.failed_pages:
                report.blank_pages.append(number)
        return pages
    finally:
        if isinstance(source, (str, os.PathLike)):
            stream.close()


def _render_rows(header: Optional[List[str]], rows: List[List[str]], continued: bool = False) -> str:
    """Same self-describing row format as CSV/DOCX tables; the header line is not repeated for a continued table."""
    lines = [] if (continued or not header) else [" | ".join(header)]
    lines += [t for t in (_row_text(header, r) for r in rows) if t]
    return "\n".join(lines)


def _table_entry(page: int, kind: str, header: Optional[List[str]], rows: List[List[str]], continued: bool = False) -> Dict[str, Any]:
    return {"page": page, "kind": kind, "rows": len(rows), "columns": len(rows[0]) if rows else 0, "header": bool(header), "continued": continued}


def _structure_tables(source, cfg, report: ExtractionReport, texts: List[Optional[str]], methods: List[str]) -> None:
    """Replace the flat text of pages that contain a recognised table by: text above, the table as rows, text below.
    Never fatal: any failure leaves the plain text of the affected page(s) and adds a warning."""
    try:
        import pdfplumber
    except Exception:
        report.warn(None, "TABLES_UNAVAILABLE", "Table extraction is enabled but pdfplumber is not installed; tables stay plain text.")
        return
    stream = _stream(source)
    budget = pdftables.Budget(cfg.PDF_TABLE_TIMEOUT_SECONDS)
    per_page: List[List[pdftables.FoundTable]] = []
    layouts: Dict[int, list] = {}
    try:
        with pdfplumber.open(cast(Any, stream)) as pdf:
            if len(pdf.pages) != len(texts):
                report.warn(None, "TABLES_FAILED", "Table extraction skipped: the PDF's page count is inconsistent between parsers.")
                return
            for i, page in enumerate(pdf.pages):
                if budget.exhausted():
                    report.warn(i + 1, "TABLES_TIME_BUDGET", f"Table extraction stopped after {cfg.PDF_TABLE_TIMEOUT_SECONDS:.0f} s; "
                                "this and later pages keep plain text.")
                    per_page += [[] for _ in range(len(pdf.pages) - i)]
                    break
                tables: List[pdftables.FoundTable] = []
                if texts[i] and methods[i] == "native":
                    try:
                        tables = pdftables.find_page_tables(page, cfg.PDF_TABLES)
                        if tables:
                            layouts[i] = [(k, (normalize(textclean.clean_chars(v.extract_text() or "")) if k == "text" else v))
                                          for k, v in pdftables.split_page(page, tables)]
                    except Exception:
                        tables = []
                        report.warn(i + 1, "TABLES_FAILED", "Table detection failed on this page; it keeps plain text.")
                per_page.append(tables)
                try:
                    page.flush_cache()
                except Exception:
                    pass
    except Exception:
        report.warn(None, "TABLES_FAILED", "The PDF could not be analysed for tables; all pages keep plain text.")
        return
    finally:
        if isinstance(source, (str, os.PathLike)):
            stream.close()

    pdftables.link_continuations(per_page, _usable_header)
    for i, parts in layouts.items():
        out = []
        for kind, value in parts:
            if kind == "text":
                if value:
                    out.append(value)
            else:
                out.append(_render_rows(value.header, value.rows, value.continued))
                report.tables.append(_table_entry(i + 1, value.kind, value.header, value.rows, value.continued))
        texts[i] = "\n\n".join(p for p in out if p)
        METRICS.inc("pdf_tables", kind=",".join(sorted({v.kind for k, v in parts if k == "table"})))


def _structure_ocr_tables(text: str, page: int, report: ExtractionReport) -> str:
    """OCR rows whose cells are separated by wide gaps (' | ') and that line up across >= 3 consecutive lines are
    rendered with the same row format. The relationships are inferred from positions, so the page is flagged."""
    blocks = pdftables.ocr_table_blocks(text)
    if not blocks:
        return text
    lines = text.split("\n")
    for start, end, rows in reversed(blocks):
        header = _usable_header(rows[0]) if len(rows) > 1 else None
        body = rows[1:] if header else rows
        lines[start:end] = _render_rows(header, body).split("\n")
        report.tables.append(_table_entry(page, "ocr_layout", header, body))
    report.warn(page, "TABLE_FROM_OCR_LAYOUT", "A table on this page was structured from OCR text positions; check the values against the original.")
    return "\n".join(lines)


def _ocr_pages(source, cfg, report, texts, methods, confs, images, candidates) -> None:
    report.ocr_engine = ocr.get_engine().name
    try:
        pdf = ocr.open_pdfium(_pdfium_source(source))
    except Exception:
        for i in candidates:
            report.failed_pages.append(i + 1)
            report.warn(i + 1, "OCR_RENDER_FAILED", "The page could not be rendered for OCR.")
        return
    started, done, scores = time.monotonic(), 0, []
    try:
        for i in candidates:
            number = i + 1
            if done >= cfg.OCR_MAX_PAGES_PER_DOCUMENT:
                report.unprocessed_pages.append(number)
                report.warn(number, "OCR_SKIPPED_PAGE_LIMIT", f"OCR is limited to {cfg.OCR_MAX_PAGES_PER_DOCUMENT} pages per document.")
                continue
            if time.monotonic() - started > cfg.OCR_DOCUMENT_TIMEOUT_SECONDS:
                report.unprocessed_pages.append(number)
                report.warn(number, "OCR_SKIPPED_TIME_BUDGET", "The OCR time budget for this document was used up.")
                continue
            page_started = time.perf_counter()
            try:
                img = ocr.render_pdf_page(pdf, i, cfg.OCR_RENDER_DPI, cfg.OCR_MAX_PIXELS)
                if _looks_blank(img):
                    continue  # a genuinely blank page: nothing to read, nothing lost
                result = ocr.recognize_with_timeout(img, cfg.OCR_PAGE_TIMEOUT_SECONDS)
                METRICS.observe_ms("ocr_page", round((time.perf_counter() - page_started) * 1000, 2), outcome="ok")
            except ocr.OcrTimeout:
                METRICS.inc("ocr_page_failures", reason="timeout")
                report.failed_pages.append(number)
                report.warn(number, "OCR_TIMEOUT", f"OCR of this page took longer than {cfg.OCR_PAGE_TIMEOUT_SECONDS:.0f} s.")
                continue
            except Exception:
                METRICS.inc("ocr_page_failures", reason="error")
                log.warning("ocr_page_failed", extra={"page": number})
                report.failed_pages.append(number)
                report.warn(number, "OCR_FAILED", "OCR failed for this page.")
                continue
            finally:
                done += 1
            text = normalize(textclean.clean_chars(result.text))
            if cfg.PDF_TABLES != "off":
                text = _structure_ocr_tables(text, number, report)
            if len(text) > len(texts[i] or ""):
                texts[i], methods[i], confs[i] = text, "ocr", result.confidence
                if result.confidence is not None:
                    scores.append(result.confidence)
                    if result.confidence < cfg.OCR_MIN_CONFIDENCE:
                        report.warn(number, "OCR_LOW_CONFIDENCE",
                                    f"OCR confidence for this page is low ({result.confidence:.2f}); the text may contain errors.")
    finally:
        try:
            pdf.close()
        except Exception:
            pass
    if scores:
        report.ocr_confidence_mean = sum(scores) / len(scores)


# ---------------------------------------------------------------- images

def _parse_image(source: Source, cfg, report: ExtractionReport) -> List[Page]:
    if not cfg.OCR_ENABLED:
        raise errors.unsupported("Images cannot be searched yet: no OCR is configured.")
    from PIL import Image, ImageOps

    previous = Image.MAX_IMAGE_PIXELS
    Image.MAX_IMAGE_PIXELS = cfg.IMAGE_MAX_PIXELS  # decompression-bomb guard (PIL raises above 2x this)
    stream = _stream(source)
    try:
        try:
            with Image.open(stream) as probe:
                probe.verify()  # structural check without decoding all pixels
            stream.seek(0)
            img = Image.open(stream)
            if img.width * img.height > cfg.IMAGE_MAX_PIXELS:
                raise errors.ServiceError(errors.DOCUMENT_TOO_LARGE, "The image has too many pixels to process.", 413)
            img = ImageOps.exif_transpose(img).convert("RGB")
        except errors.ServiceError:
            raise
        except Exception as exc:  # malformed, truncated or decompression-bomb images
            raise errors.parse_failed() from exc
        report.total_pages = 1
        report.ocr_engine = ocr.get_engine().name
        if _looks_blank(img):
            report.blank_pages.append(1)
            return []
        try:
            result = ocr.recognize_with_timeout(img, cfg.OCR_PAGE_TIMEOUT_SECONDS)
        except Exception as exc:
            log.warning("ocr_image_failed")
            report.failed_pages.append(1)
            report.warn(1, "OCR_TIMEOUT" if isinstance(exc, ocr.OcrTimeout) else "OCR_FAILED", "OCR failed for this image.")
            return []
        text = normalize(textclean.clean_chars(result.text))
        if text and cfg.PDF_TABLES != 'off':
            text = _structure_ocr_tables(text, 1, report)
        if not text:
            report.blank_pages.append(1)
            return []
        report.ocr_pages = 1
        report.ocr_confidence_mean = result.confidence
        if result.confidence is not None and result.confidence < cfg.OCR_MIN_CONFIDENCE:
            report.warn(1, "OCR_LOW_CONFIDENCE", f"OCR confidence is low ({result.confidence:.2f}); the text may contain errors.")
        return [Page(None, text, "ocr", result.confidence)]
    finally:
        Image.MAX_IMAGE_PIXELS = previous
        if isinstance(source, (str, os.PathLike)):
            stream.close()


# ---------------------------------------------------------------- entry points

def parse_document(source: Source, mime_type: str, cfg=None) -> ParsedDocument:
    if cfg is None:
        from rag_core.config import settings as cfg
    if mime_type in IMAGE_TYPES:
        report = ExtractionReport(source_type="image")
        try:
            pages = _parse_image(source, cfg, report)
        except errors.ServiceError:
            raise
        except Exception as exc:
            raise errors.parse_failed() from exc
    else:
        report = ExtractionReport(source_type={PDF: "pdf", DOCX: "docx", CSV: "csv", MARKDOWN: "markdown"}.get(mime_type, "text"))
        try:
            if mime_type == PDF:
                pages = _parse_pdf(source, cfg, report)
            elif mime_type == DOCX:
                _check_zip_safety(source, cfg)
                raw, wanted = _docx_text(source)
                text = normalize(textclean.clean_chars(raw))
                pages = [Page(None, text, headings=locate_headings(text, wanted))] if text else []
                report.native_pages = 1 if pages else 0
            elif mime_type == CSV:
                pages = [Page(None, normalize(textclean.clean_chars(_csv_text(_decode(_read_bytes(source))))))]
                report.native_pages = 1
            elif mime_type in TEXT_TYPES:
                text = normalize(textclean.clean_chars(_decode(_read_bytes(source))))
                if cfg.TEXT_CLEANUP:
                    text, joins = textclean.dehyphenate(text)
                    report.cleanup = {"hyphen_joins": joins}
                head = markdown_headings(text) if mime_type == MARKDOWN else []
                pages = [Page(None, text, headings=head)]
                report.native_pages = 1
            else:
                raise errors.unsupported("This file type is not supported.")
        except errors.ServiceError:
            raise
        except Exception as exc:  # malformed files: never retry, never leak parser internals
            raise errors.parse_failed() from exc
    pages = [p for p in pages if p.text]
    if not pages:
        if report.unprocessed_pages and not cfg.OCR_ENABLED:
            raise errors.ServiceError(
                errors.NO_EXTRACTABLE_TEXT,
                "No extractable text was found: the document appears to be scanned or image-only and OCR is not enabled.",
                422,
            )
        raise errors.no_text()
    report.total_pages = report.total_pages if report.total_pages is not None else (len(pages) if report.source_type == "pdf" else None)
    return ParsedDocument(pages, report)


def parse(data: Source, mime_type: str) -> List[Page]:
    """Original interface: pages only (the extraction report is available from ``parse_document``)."""
    return parse_document(data, mime_type).pages
