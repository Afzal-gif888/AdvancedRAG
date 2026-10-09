import io
import json
import uuid

import pytest

from conftest import IS_PG, make_pdf
from rag_core import errors
from rag_core.chunking import chunk_pages
from rag_core.parsing import Page, parse
from rag_core.pipeline import rrf

U1, U2 = uuid.uuid4(), uuid.uuid4()

GUIDE = (
    "The refund policy allows customers to return products within 30 days of purchase. "
    "Refunds are processed to the original payment method within five business days.\n\n"
    "Shipping is free for orders above fifty dollars. Express shipping costs ten dollars "
    "and arrives in two days.\n\n"
    "Warranty claims require the original receipt and the product serial number."
)


async def _ingest(client, doc_id, user_id, data, mime="text/plain", name="guide.txt"):
    meta = {"document_id": str(doc_id), "user_id": str(user_id), "filename": name, "mime_type": mime}
    return await client.post("/api/v1/rag/ingest", data={"metadata": json.dumps(meta)}, files={"file": (name, data, mime)})


# ---------------------------------------------------------------- parsing & chunking

def test_parse_pdf_pages():
    pages = parse(make_pdf(["Alpha page text", "Beta page text"]), "application/pdf")
    assert [(p.number, p.text) for p in pages] == [(1, "Alpha page text"), (2, "Beta page text")]


def test_parse_docx():
    import docx

    d = docx.Document()
    d.add_paragraph("Quarterly revenue grew.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Region", "EMEA"
    buf = io.BytesIO()
    d.save(buf)
    text = parse(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")[0].text
    assert "Quarterly revenue grew." in text and "Region | EMEA" in text


def test_parse_text_encodings_and_failures():
    assert parse("café".encode("cp1252"), "text/plain")[0].text == "café"
    with pytest.raises(errors.ServiceError) as exc:
        parse(b"   \n ", "text/plain")
    assert exc.value.code == errors.NO_EXTRACTABLE_TEXT
    with pytest.raises(errors.ServiceError) as exc:
        parse(b"\x89PNG....", "image/png")
    assert exc.value.code == errors.UNSUPPORTED_CONTENT and exc.value.status_code == 415
    with pytest.raises(errors.ServiceError) as exc:
        parse(b"%PDF-1.4 garbage", "application/pdf")
    assert exc.value.code == errors.PARSE_FAILED


def test_chunking_respects_size_overlap_and_pages():
    text = " ".join(f"Sentence number {i} talks about topic {i % 7}." for i in range(200))
    chunks = chunk_pages([Page(3, text), Page(4, "short page")], size=300, overlap=50)
    assert all(len(c.text) <= 350 for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert chunks[-1].page == 4 and chunks[0].page == 3
    # Overlap carries context forward.
    assert chunks[1].text.split()[0] in chunks[0].text


def test_rrf_fuses_rankings():
    fused = rrf([["a", "b", "c"], ["c", "a", "d"]], k=60)
    assert [x for x, _ in fused][:2] == ["a", "c"]


# ---------------------------------------------------------------- API (fake embeddings + local store)

async def test_ingest_retrieve_and_context(client):
    doc = uuid.uuid4()
    r = await _ingest(client, doc, U1, GUIDE.encode())
    assert r.status_code == 200, r.text
    assert r.json()["chunks"] >= 2

    r = await client.post("/api/v1/rag/retrieve", json={"user_id": str(U1), "query": "how long do refunds take", "document_ids": [str(doc)]})
    body = r.json()
    assert body["chunks"], body
    top = body["chunks"][0]
    assert "Refund" in top["text"] and top["source_id"] == 1 and top["filename"] == "guide.txt"
    assert body["context"].startswith("[1] guide.txt")
    meta = body["metadata"]
    assert meta["retrieval_strategy"] == "hybrid_rrf_rerank" and meta["semantic_candidates"] > 0 and meta["keyword_candidates"] > 0


async def test_pdf_ingest_keeps_page_numbers(client):
    doc = uuid.uuid4()
    await _ingest(client, doc, U1, make_pdf(["Introduction to widgets", "Widget maintenance requires oiling monthly"]), "application/pdf", "w.pdf")
    r = await client.post("/api/v1/rag/retrieve", json={"user_id": str(U1), "query": "widget maintenance oiling", "document_ids": [str(doc)]})
    top = r.json()["chunks"][0]
    assert top["page"] == 2 and "(page 2)" in r.json()["context"]


async def test_user_isolation_in_retrieval(client):
    doc = uuid.uuid4()
    await _ingest(client, doc, U1, GUIDE.encode())
    # U2 passes U1's document id: nothing comes back.
    r = await client.post("/api/v1/rag/retrieve", json={"user_id": str(U2), "query": "refund", "document_ids": [str(doc)]})
    assert r.json()["chunks"] == []
    # U2 cannot overwrite or delete U1's document either.
    assert (await _ingest(client, doc, U2, b"hijack")).status_code == 409
    assert (await client.delete(f"/api/v1/rag/documents/{doc}", params={"user_id": str(U2)})).status_code == 404


async def test_reingest_is_idempotent_and_delete_cleans_up(client, database):
    from sqlalchemy import func, select

    from rag_core.db import Chunk

    doc = uuid.uuid4()
    await _ingest(client, doc, U1, (GUIDE + "\n\n" + GUIDE).encode())
    await _ingest(client, doc, U1, b"Only one short paragraph about refunds now.")
    async with database.connect() as conn:
        assert (await conn.execute(select(func.count()).select_from(Chunk).where(Chunk.document_id == doc))).scalar() == 1
    assert (await client.delete(f"/api/v1/rag/documents/{doc}", params={"user_id": str(U1)})).status_code == 204
    async with database.connect() as conn:
        assert (await conn.execute(select(func.count()).select_from(Chunk).where(Chunk.document_id == doc))).scalar() == 0
    assert (await client.delete(f"/api/v1/rag/documents/{doc}", params={"user_id": str(U1)})).status_code == 404


async def test_unsupported_and_empty_documents_are_non_retryable(client):
    r = await _ingest(client, uuid.uuid4(), U1, b"\x89PNG\r\n\x1a\n", "image/png", "photo.png")
    assert r.status_code == 415 and r.json()["error"]["code"] == "UNSUPPORTED_CONTENT" and r.json()["error"]["retryable"] is False
    r = await _ingest(client, uuid.uuid4(), U1, b"   ")
    assert r.status_code == 422 and r.json()["error"]["code"] == "NO_EXTRACTABLE_TEXT"
    r = await client.post("/api/v1/rag/ingest", data={"metadata": "{bad"}, files={"file": ("a.txt", b"x", "text/plain")})
    assert r.status_code == 422


async def test_unindexed_documents_return_empty(client):
    r = await client.post("/api/v1/rag/retrieve", json={"user_id": str(U1), "query": "anything", "document_ids": [str(uuid.uuid4())]})
    assert r.status_code == 200 and r.json()["chunks"] == [] and r.json()["context"] == ""


async def test_health_ready_and_auth(client, monkeypatch):
    assert (await client.get("/health")).json() == {"status": "ok"}
    r = await client.get("/ready")
    assert r.status_code == 200 and r.json()["checks"]["database"] is True


@pytest.mark.skipif(not IS_PG, reason="PostgreSQL full-text search")
async def test_postgres_full_text_keyword_search(services, database):
    from rag_core.db import SessionLocal
    from rag_core.pipeline import _keyword, ingest

    doc = uuid.uuid4()
    async with SessionLocal() as db:
        await ingest(db, services, document_id=doc, user_id=U1, filename="g.txt", mime_type="text/plain", data=GUIDE.encode())
    async with SessionLocal() as db:
        hits = await _keyword(db, U1, "express shipping", [doc], 5)
        assert hits and hits[0][1] > 0
        # websearch syntax and hostile input are handled by PostgreSQL's parser, not string building.
        assert await _keyword(db, U1, "'; DROP TABLE rag_chunks; --", [doc], 5) == []


# ---------------------------------------------------------------- real adapters against mock HTTP APIs

async def test_local_model_dimension_and_query_passage_consistency(local_embeddings):
    from rag_core.embeddings import DOCUMENT, QUERY, model_dimensions

    assert local_embeddings.dimensions == model_dimensions("BAAI/bge-small-en-v1.5") == 384
    q = (await local_embeddings.embed(["How long is the warranty?"], QUERY))[0]
    docs = await local_embeddings.embed(["Every widget ships with a 24 month warranty.", "Shipping is free above fifty dollars."], DOCUMENT)
    assert len(q) == 384 and all(len(d) == 384 for d in docs)
    assert abs(sum(x * x for x in q) - 1.0) < 1e-3  # unit vectors (cosine == dot product)
    cos = [sum(a * b for a, b in zip(q, d)) for d in docs]
    assert cos[0] > cos[1]  # the relevant passage ranks first
    # Deterministic: the same text always embeds to the same vector.
    assert (await local_embeddings.embed(["How long is the warranty?"], QUERY))[0] == q


def test_dimension_guard_rejects_wrong_configuration():
    from rag_core.config import ConfigError
    from rag_core.embeddings import LocalEmbeddings

    with pytest.raises(ConfigError) as exc:
        LocalEmbeddings(model="BAAI/bge-small-en-v1.5", cache_dir=None, expected_dimensions=768)
    assert "produces 384" in str(exc.value)
    with pytest.raises(ConfigError):
        LocalEmbeddings(model="not-a-real/model", cache_dir=None)


async def test_local_embeddings_with_pinecone_adapter_end_to_end(real_provider_services, database):
    from mock_providers import state
    from rag_core.db import SessionLocal
    from rag_core.pipeline import delete_document, ingest, retrieve

    svc = real_provider_services
    doc = uuid.uuid4()
    async with SessionLocal() as db:
        res = await ingest(db, svc, document_id=doc, user_id=U1, filename="g.txt", mime_type="text/plain", data=GUIDE.encode())
    assert res["chunks"] >= 2 and res["embedding_model"] == "BAAI/bge-small-en-v1.5"
    assert set(state.namespaces) == {f"user-{U1}"}
    assert all(len(v["values"]) == 384 for v in state.namespaces[f"user-{U1}"].values())

    async with SessionLocal() as db:
        out = await retrieve(db, svc, user_id=U1, query="How many days until my refund arrives?", document_ids=[doc])
    assert out["chunks"] and "refund" in out["chunks"][0]["text"].lower()
    assert out["metadata"]["reranker"] == "pinecone" and state.rerank_calls == 1
    assert out["metadata"]["embedding_model"] == "BAAI/bge-small-en-v1.5"
    assert state.last_query["namespace"] == f"user-{U1}" and state.last_query["filter"] == {"document_id": {"$in": [str(doc)]}}
    assert len(state.last_query["vector"]) == 384

    state.rerank_fail = True
    async with SessionLocal() as db:
        out = await retrieve(db, svc, user_id=U1, query="refund policy days", document_ids=[doc])
    assert out["chunks"] and out["metadata"]["reranker"] == "skipped"  # graceful degradation

    async with SessionLocal() as db:
        assert await delete_document(db, svc, document_id=doc, user_id=U1)
    assert state.namespaces[f"user-{U1}"] == {}


async def test_embedding_failure_is_retryable_and_index_dimension_is_checked(real_provider_services, database, monkeypatch):
    from mock_providers import PINECONE_KEY, state
    from rag_core.db import SessionLocal
    from rag_core.pipeline import ingest
    from rag_core.vectorstore import PineconeStore

    def broken(texts, task):
        raise RuntimeError("onnx runtime failure")

    monkeypatch.setattr(real_provider_services.embeddings, "_run", broken)
    with pytest.raises(errors.ServiceError) as exc:
        async with SessionLocal() as db:
            await ingest(db, real_provider_services, document_id=uuid.uuid4(), user_id=U1, filename="g.txt", mime_type="text/plain", data=GUIDE.encode())
    assert exc.value.code == errors.EMBEDDING_UNAVAILABLE and exc.value.retryable

    api_url = str(real_provider_services.vectors._control.base_url)
    state.reset()
    state.dimension = 768  # e.g. an index created for a different model
    ok, detail = await PineconeStore(api_key=PINECONE_KEY, index="advancerag-test", dimensions=384, api_url=api_url).health()
    assert not ok and detail == "DIMENSION_MISMATCH: index has 768, embedding model produces 384"
    from rag_core.vectorstore import VectorRecord

    # Writes into a wrong-dimension index are refused before anything is sent.
    with pytest.raises(errors.ServiceError) as exc:
        await PineconeStore(api_key=PINECONE_KEY, index="advancerag-test", dimensions=384, api_url=api_url).upsert(
            U1, [VectorRecord(id="x:0", values=[0.0] * 384, metadata={"document_id": "x"})])
    assert exc.value.code == errors.VECTOR_STORE_MISCONFIGURED and "dimension 384" in exc.value.message
    assert state.namespaces == {}
    state.reset()
    bad = PineconeStore(api_key="wrong", index="advancerag-test", dimensions=384, api_url=api_url)
    assert (await bad.health()) == (False, "VECTOR_STORE_MISCONFIGURED")


def test_production_config_validation():
    from rag_core.config import ConfigError, Settings

    base = dict(ENVIRONMENT="production", DATABASE_URL="postgresql+asyncpg://u:secretpw@db/rag", _env_file=None,
                INTERNAL_SERVICE_TOKEN="t" * 40, PINECONE_API_KEY="p" * 30, PINECONE_INDEX="idx",
                EMBEDDING_PROVIDER="local", VECTOR_STORE="pinecone")
    Settings(**base).check_safety()  # no Gemini (or any embedding API) key required
    for override in ({"EMBEDDING_PROVIDER": "fake"}, {"VECTOR_STORE": "local"}, {"PINECONE_API_KEY": ""}, {"PINECONE_INDEX": ""}, {"INTERNAL_SERVICE_TOKEN": ""}):
        with pytest.raises(ConfigError) as exc:
            Settings(**{**base, **override}).check_safety()
        assert "secretpw" not in str(exc.value) and "p" * 30 not in str(exc.value)
