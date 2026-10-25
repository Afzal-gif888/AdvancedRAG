"""AdvancedRAG retrieval benchmark against the RUNNING stack (the frozen RAG baseline).

    python scripts/rag_benchmark.py        # from the repository root, with the dev stack running

What it does, all under a fresh benchmark user (its own Pinecone namespace):
  1. ingests a fixed mixed-format corpus (PDF, 2x DOCX, TXT, MD, CSV, distractor) through RAG Core
  2. verifies Pinecone: vector count, 384-d unit vectors, metadata, vector == embedding(chunk text)
  3. scores retrieval per stage (semantic / keyword / RRF / final) on 20 questions + 6 unanswerable
  4. runs 9 end-to-end questions through Agent Core -> LLM service -> FreeLLMAPI
  5. deletes everything it created and checks the namespace is empty

Needs: rag-core/.env (Pinecone key, internal token), services on :8001/:8002/:8003.
Frozen baseline (2026-10-08): final hit@1 19/20, MRR 0.975; E2E 7/7 + 2/2. See docs/rag-architecture.md.
"""
import asyncio
import json
import math
import os
import re
import sys
import time
import uuid
from pathlib import Path

ROOT = Path.cwd()
os.environ["ADVANCERAG_ENV_FILE"] = str(ROOT / "rag-core" / ".env")
sys.path.insert(0, str(ROOT / "rag-core"))

import io  # noqa: E402

import httpx  # noqa: E402

from rag_core import db as rag_db  # noqa: E402
from rag_core import pipeline  # noqa: E402
from rag_core.config import settings  # noqa: E402
from rag_core.db import Chunk  # noqa: E402
from rag_core.embeddings import QUERY, LocalEmbeddings  # noqa: E402
from rag_core.vectorstore import build_vector_store, namespace_for  # noqa: E402
from sqlalchemy import func, select  # noqa: E402


PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_pdf(pages):
    """Valid PDF, several text lines per page."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objs, kids = [], []
    font_id = 3 + 2 * len(pages)
    for i, lines in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        kids.append(f"{page_id} 0 R")
        ops = "BT /F1 10 Tf 14 TL 50 760 Td " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET"
        stream = ops.encode("latin-1")
        objs.append((page_id, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()))
        objs.append((content_id, b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"))
    objs = [(1, b"<< /Type /Catalog /Pages 2 0 R >>"), (2, f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode())] + objs
    objs.append((font_id, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"))
    objs.sort()
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for oid, body in objs:
        offsets[oid] = len(out)
        out += f"{oid} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for oid, _ in objs:
        out += f"{offsets[oid]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


FILLER = ("Northwind Labs was founded in a small workshop and grew through careful engineering, patient "
          "customer work and a culture of writing things down. ")


def corpus():
    pdf = make_pdf([
        ["Northwind Labs Employee Handbook", "Section 1: Remote work",
         "Employees may work remotely up to three days per week, provided their manager approves",
         "the schedule in advance. Core collaboration hours are 10:00 to 15:00 CET.",
         FILLER, FILLER, FILLER],
        ["Section 2: Travel and expenses",
         "Hotel stays are reimbursed up to 180 euros per night in all European cities.",
         "Flights must be booked in economy class unless the flight time exceeds 9 hours,",
         "in which case premium economy is allowed. Receipts must be submitted within 30 days.",
         FILLER, FILLER],
        ["Section 3: Information security",
         "Laptops must use full-disk encryption. Security incidents must be reported to",
         "security@northwind.example within 24 hours of discovery. Lost devices are wiped remotely.",
         FILLER, FILLER],
        ["Section 4: Learning and development",
         "Every employee has an annual training budget of 1,500 euros for courses, books and conferences.",
         "Unused budget does not roll over to the next year.", FILLER, FILLER],
        ["Section 5: Parental leave",
         "Parents receive 16 weeks of fully paid parental leave, which can be split into two blocks",
         "within the first two years after birth or adoption.", FILLER, FILLER],
        ["Section 6: Equipment",
         "New hires choose a laptop from the approved list and receive a 400 euro home-office allowance",
         "for a desk, chair or monitor.", FILLER, FILLER],
        ["Section 7: Public holidays",
         "The office follows the public holidays of the contract country plus December 24 and 31.",
         FILLER, FILLER],
        ["Section 8: Code of conduct",
         "Gifts from suppliers above 50 euros must be declared to the compliance team.", FILLER, FILLER],
    ])

    import docx
    d = docx.Document()
    d.add_heading("Aurora X2 Product Specification", 0)
    d.add_heading("Overview", 1)
    d.add_paragraph("The Aurora X2 is a wireless air-quality sensor for offices and warehouses. It measures CO2, "
                    "particulate matter and humidity and reports every five minutes over LoRaWAN.")
    d.add_heading("Technical specifications", 1)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Property", "Value"
    for k, v in [("Battery life", "14 months"), ("Operating temperature", "-20 to 60 degrees Celsius"),
                 ("Weight", "85 grams"), ("Ingress protection", "IP65")]:
        row = t.add_row()
        row.cells[0].text, row.cells[1].text = k, v
    d.add_heading("Warranty", 1)
    d.add_paragraph("The warranty covers manufacturing defects for 24 months from delivery. Water damage beyond "
                    "the IP65 rating is not covered.")
    buf = io.BytesIO()
    d.save(buf)

    txt = ("Product meeting notes, 12 June\n\nAttendees: Priya, Jonas, Mei.\n\nDecision: the Q3 launch of the Aurora X2 "
           "was moved to September 18 because the external certification lab was delayed by three weeks.\n\n"
           "Action items: Jonas updates the launch plan; Mei informs distributors.\n")

    md = ("# Operations runbook\n\n## Database failover\n\nIf the primary database is unavailable, promote the replica by "
          "running `pgctl promote` on the standby host, then update the connection string in the secrets store. "
          "The recovery time objective (RTO) is 15 minutes.\n\n## Log rotation\n\nLogs are rotated daily and kept for 14 days.\n")

    d1 = docx.Document()
    d1.add_heading("Aurora X1 Product Specification", 0)
    d1.add_paragraph("The Aurora X1 is the previous-generation indoor sensor. It measures CO2 and temperature only.")
    t1 = d1.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text, t1.rows[0].cells[1].text = "Property", "Value"
    for k, v in [("Battery life", "9 months"), ("Operating temperature", "-10 to 50 degrees Celsius"), ("Weight", "110 grams")]:
        row = t1.add_row()
        row.cells[0].text, row.cells[1].text = k, v
    d1.add_heading("Warranty", 1)
    d1.add_paragraph("The Aurora X1 warranty covers 12 months from delivery.")
    buf1 = io.BytesIO()
    d1.save(buf1)

    rows = ["product,country,price_eur,sku"]
    regions = ["Germany", "France", "Italy", "Spain", "Netherlands", "Poland", "Sweden", "Austria", "Belgium", "Portugal",
               "Denmark", "Finland", "Ireland", "Czechia", "Hungary", "Greece", "Romania", "Slovakia", "Slovenia", "Croatia",
               "Norway", "Switzerland", "Estonia", "Latvia", "Lithuania", "Luxembourg", "Malta", "Cyprus", "Bulgaria", "Iceland"]
    products = ["Aurora X1", "Aurora X2", "Borealis Hub", "Nimbus Gateway", "Cirrus Mount", "Stratus Cable", "Zephyr Relay", "Halo Display"]
    prices = {"Aurora X1": 199, "Aurora X2": 249, "Borealis Hub": 129, "Nimbus Gateway": 349, "Cirrus Mount": 19,
              "Stratus Cable": 9, "Zephyr Relay": 89, "Halo Display": 159}
    for p in products:
        for i, r in enumerate(regions):
            code = "".join(w[0] for w in p.split()) + p.split()[-1][-1]
            rows.append(f"{p},{r},{prices[p] + i * 3},{code}-{r[:2].upper()}{i}")
    csv = "\n".join(rows) + "\n"

    history = "\n\n".join(
        f"Chapter {i}. " + FILLER * 4 + f"In year {2010 + i} the team opened an office in city number {i} and hired {i * 3} people."
        for i in range(1, 41))

    return [
        ("employee_handbook.pdf", PDF, pdf),
        ("aurora_x2_spec.docx", DOCX, buf.getvalue()),
        ("aurora_x1_spec.docx", DOCX, buf1.getvalue()),
        ("meeting_notes.txt", "text/plain", txt.encode()),
        ("runbook.md", "text/markdown", md.encode()),
        ("price_list.csv", "text/csv", csv.encode()),
        ("company_history.txt", "text/plain", history.encode()),
    ]


# (question, filename that must be cited, phrases that must appear in the chunk)
QUESTIONS = [
    ("How many days a week can I work from home?", "employee_handbook.pdf", ["three days"]),
    ("What is the maximum hotel cost per night on a business trip?", "employee_handbook.pdf", ["180 euros"]),
    ("When am I allowed to fly premium economy?", "employee_handbook.pdf", ["9 hours"]),
    ("How quickly do I have to report a security incident?", "employee_handbook.pdf", ["24 hours"]),
    ("How long does the Aurora X2 battery last?", "aurora_x2_spec.docx", ["14 months", "Aurora X2"]),
    ("What temperature range can the sensor operate in?", "aurora_x2_spec.docx", ["-20 to 60"]),
    ("How long is the warranty on the Aurora X2?", "aurora_x2_spec.docx", ["24 months"]),
    ("Why was the Q3 launch delayed?", "meeting_notes.txt", ["certification lab"]),
    ("How do I promote the standby during a database failover?", "runbook.md", ["pgctl promote"]),
    ("What is the RTO for database failover?", "runbook.md", ["15 minutes"]),
    ("What does the Aurora X2 cost in Germany?", "price_list.csv", ["Aurora X2", "Germany", "249"]),
    ("What is the price of the Halo Display in Portugal?", "price_list.csv", ["Halo Display", "Portugal"]),
    ("Which product has the SKU ZRy-SL18?", "price_list.csv", ["ZRy-SL18"]),
    ("How much does the Stratus Cable cost in Iceland?", "price_list.csv", ["Stratus Cable", "Iceland"]),
    ("Battery life of the older X1 sensor?", "aurora_x1_spec.docx", ["9 months"]),
    ("Can I do my job from my apartment most of the week?", "employee_handbook.pdf", ["three days"]),
    ("Is there money for courses and conferences?", "employee_handbook.pdf", ["1,500 euros"]),
    ("How much time off do new parents get?", "employee_handbook.pdf", ["16 weeks"]),
    ("Do I have to declare presents from vendors?", "employee_handbook.pdf", ["50 euros"]),
    ("Is there an allowance for a desk or chair at home?", "employee_handbook.pdf", ["400 euro"]),
]
NEGATIVE = ["What is the capital of Australia?", "Who won the 1998 football world cup?",
            "What dental insurance do employees get?", "What is the company pension contribution?",
            "How many vacation days do I get per year?", "What is the Bluetooth range of the Aurora X2?"]
BENCH_USER = uuid.uuid5(uuid.NAMESPACE_URL, "advancerag/rag-benchmark")  # namespace user-<this>
TOKEN = settings.internal_token
AUTH = {"Authorization": f"Bearer {TOKEN}"}
RAG, AGENT = "http://127.0.0.1:8002", "http://127.0.0.1:8001"

E2E = [
    ("What does the Aurora X2 cost in Germany?", [], ["249"], "price_list.csv"),
    ("How much is the Halo Display in Portugal?", [], ["186"], "price_list.csv"),
    ("How long is the warranty on the Aurora X2?", [], ["24 months", "24-month", "two years"], "aurora_x2_spec.docx"),
    ("What's the battery life of the older X1 sensor?", [], ["9 months", "nine months"], "aurora_x1_spec.docx"),
    ("Can I do my job from my apartment most of the week?", [], ["three days", "3 days"], "employee_handbook.pdf"),
    ("What is the RTO for database failover?", [], ["15 minutes", "15-minute"], "runbook.md"),
    ("And what about its operating temperature?",
     [{"role": "user", "content": "How long does the Aurora X2 battery last?"},
      {"role": "assistant", "content": "The Aurora X2 battery lasts 14 months [1]."}],
     ["-20", "minus 20"], "aurora_x2_spec.docx"),
]
UNANSWERABLE = ["What is the Bluetooth range of the Aurora X2?", "What dental insurance do employees get?"]
# The system's own abstain sentence (agent-core/agent_core/prompts.py ABSTAIN_MESSAGE) is the primary signal; the
# regex below only catches free-form refusals from answers that did not go through the abstain path. Before this
# line was added the check missed the abstain sentence and would have failed a correct run.
ABSTAIN_SENTENCE = "i couldn't find sufficient evidence"
REFUSAL = re.compile(r"(not (mention|specif|includ|contain|cover|state|provid|list|found|available)|doesn.t (mention|cover|specify|include|contain|say|list)|no (information|mention|details)|don.t (mention|cover|specify|include|contain)|do not (mention|cover|specify|include|contain|list))", re.I)


def norm(s):
    return re.sub(r"[   ]", " ", s).replace("‑", "-").replace("−", "-").lower()


async def pinecone_stats(store, user):
    client = await store._client()
    r = await client.post("/describe_index_stats", headers=store._headers(), json={})
    stats = r.json()
    return (stats.get("namespaces") or {}).get(namespace_for(user), {}).get("vectorCount", 0), stats.get("totalVectorCount")


async def main():
    # Fixed benchmark identity: a run interrupted before cleanup is cleaned by the next run.
    user = BENCH_USER
    store = build_vector_store(settings, 384)
    rag_db.configure_engine(settings.database_url)
    emb = LocalEmbeddings(model=settings.EMBEDDING_MODEL, cache_dir=str(ROOT / "rag-core" / settings.EMBEDDING_CACHE_DIR))
    http = httpx.AsyncClient(timeout=300)
    docs = {}  # doc_id -> filename
    checks = {}  # baseline gate: name -> passed
    print(f"index={settings.PINECONE_INDEX} model={settings.EMBEDDING_MODEL} reranker={settings.RERANKER} benchmark namespace={namespace_for(user)}")
    try:
        # ---------------------------------------------------------------- 0. remove leftovers of an interrupted run
        from rag_core.db import IndexedDocument

        async with rag_db.SessionLocal() as db:
            stale = (await db.scalars(select(IndexedDocument.document_id).where(IndexedDocument.user_id == user))).all()
        for did in stale:
            await http.delete(f"{RAG}/api/v1/rag/documents/{did}", headers=AUTH, params={"user_id": str(user)})
        if stale:
            print(f"removed {len(stale)} leftover benchmark document(s) from an interrupted run")

        # ---------------------------------------------------------------- 1. ingest through the running service
        print("\n=== INGEST (running RAG Core -> Pinecone) ===")
        expected = 0
        for name, mime, data in corpus():
            did = uuid.uuid4()
            t = time.perf_counter()
            meta = {"document_id": str(did), "user_id": str(user), "filename": name, "mime_type": mime}
            r = await http.post(f"{RAG}/api/v1/rag/ingest", headers=AUTH, data={"metadata": json.dumps(meta)}, files={"file": (name, data, mime)})
            r.raise_for_status()
            out = r.json()
            docs[did] = name
            expected += out["chunks"]
            print(f"  {name:24} chunks={out['chunks']:3} pages={out['pages']} model={out['embedding_model']} {round((time.perf_counter() - t) * 1000)} ms")

        # ---------------------------------------------------------------- 2. verify Pinecone
        print("\n=== VERIFY PINECONE ===")
        deadline = time.time() + 120
        while True:
            count, total = await pinecone_stats(store, user)
            if count == expected or time.time() > deadline:
                break
            await asyncio.sleep(3)
        print(f"  namespace vectors={count} expected={expected} (index total={total}) -> {'PASS' if count == expected else 'FAIL'}")
        async with rag_db.SessionLocal() as db:
            rows = {c.id: c for c in (await db.scalars(select(Chunk).where(Chunk.user_id == user))).all()}
        print(f"  rag_chunks rows={len(rows)} -> {'PASS' if len(rows) == expected else 'FAIL'}")
        client = await store._client()
        sample = list(rows)[:: max(1, len(rows) // 10)][:10]
        r = await client.get("/vectors/fetch", headers=store._headers(), params=[("ids", i) for i in sample] + [("namespace", namespace_for(user))])
        fetched = r.json().get("vectors", {})
        dims_ok = all(len(v["values"]) == 384 for v in fetched.values())
        unit_ok = all(abs(math.sqrt(sum(x * x for x in v["values"])) - 1) < 1e-3 for v in fetched.values())
        meta_ok = all(set(v.get("metadata", {})) <= {"document_id", "chunk_index", "page"} and v["metadata"]["document_id"] == str(rows[i].document_id)
                      and int(v["metadata"]["chunk_index"]) == rows[i].chunk_index for i, v in fetched.items())
        # Stored vector == freshly computed passage embedding of the stored chunk text (no drift / mix-up).
        recomputed = await emb.embed([rows[i].text for i in fetched], "document")
        same = all(sum(a * b for a, b in zip(fetched[i]["values"], v)) > 0.999 for i, v in zip(fetched, recomputed))
        print(f"  fetched {len(fetched)}/{len(sample)} sample vectors: dim384={dims_ok} unit-norm={unit_ok} metadata-matches-chunk={meta_ok} "
              f"no-text-in-metadata={meta_ok} vector==embedding(chunk text)={same}")
        checks["pinecone vectors == chunks"] = count == expected == len(rows)
        checks["pinecone vectors are 384-d bge embeddings of the chunk text"] = bool(fetched) and dims_ok and unit_ok and meta_ok and same
        await asyncio.sleep(5)  # serverless read-after-write freshness

        # ---------------------------------------------------------------- 3. retrieval benchmark
        print("\n=== RETRIEVAL BENCHMARK (real Pinecone) — rank of correct chunk ===")
        print(f"{'question':58} {'sem':>4} {'kw':>4} {'rrf':>4} {'final':>5}  top-file-ok  score")
        doc_ids = list(docs)

        def rank(ids, fname, phrases):
            for n, cid in enumerate(ids):
                row = rows.get(cid)
                if row and docs[row.document_id] == fname and all(p in row.text for p in phrases):
                    return n + 1
            return None

        stats = {"sem": [], "kw": [], "rrf": [], "final": []}
        latencies = []
        for q, fname, phrases in QUESTIONS:
            qv = (await emb.embed([q], QUERY))[0]
            sem = await store.query(user, qv, settings.RETRIEVAL_CANDIDATES, doc_ids)
            async with rag_db.SessionLocal() as db:
                kw = await pipeline._keyword(db, user, q, doc_ids, settings.RETRIEVAL_CANDIDATES)
            fused = pipeline.rrf([[c for c, _ in sem], [c for c, _ in kw]], settings.RRF_K)
            t = time.perf_counter()
            res = (await http.post(f"{RAG}/api/v1/rag/retrieve", headers=AUTH,
                                   json={"user_id": str(user), "query": q, "document_ids": [str(d) for d in doc_ids]})).json()
            latencies.append((time.perf_counter() - t) * 1000)
            final = [c["chunk_id"] for c in res["chunks"]]
            r_ = {"sem": rank([c for c, _ in sem], fname, phrases), "kw": rank([c for c, _ in kw], fname, phrases),
                  "rrf": rank([c for c, _ in fused], fname, phrases), "final": rank(final, fname, phrases)}
            for k, v in r_.items():
                stats[k].append(v)
            top_ok = bool(res["chunks"]) and res["chunks"][0]["filename"] == fname
            print(f"{q[:58]:58} {r_['sem'] or '-':>4} {r_['kw'] or '-':>4} {r_['rrf'] or '-':>4} {r_['final'] or '-':>5}  {str(top_ok):11}  "
                  f"{res['chunks'][0]['score'] if res['chunks'] else None}  [{res['metadata'].get('retrieval_strategy')}/{res['metadata'].get('reranker')}]")
        n = len(QUESTIONS)
        print("\n  stage   hit@1    hit@6    MRR")
        for k, ranks in stats.items():
            print(f"  {k:6} {sum(1 for x in ranks if x == 1):2}/{n}    {sum(1 for x in ranks if x and x <= 6):2}/{n}    {sum(1 / x for x in ranks if x) / n:.3f}")
        final_ranks = stats["final"]
        checks["final hit@1 >= 19/20"] = sum(1 for x in final_ranks if x == 1) >= 19
        checks["final hit@6 == 20/20"] = all(x and x <= 6 for x in final_ranks)
        latencies.sort()
        print(f"  retrieve latency (running service, incl. Pinecone query + rerank): p50={latencies[len(latencies) // 2]:.0f} ms  max={latencies[-1]:.0f} ms")
        print("\n  unanswerable queries:")
        for q in NEGATIVE:
            res = (await http.post(f"{RAG}/api/v1/rag/retrieve", headers=AUTH,
                                   json={"user_id": str(user), "query": q, "document_ids": [str(d) for d in doc_ids]})).json()
            print(f"    {q:50} top={res['chunks'][0]['filename'] if res['chunks'] else None} score={res['chunks'][0]['score'] if res['chunks'] else None}")
        other = (await http.post(f"{RAG}/api/v1/rag/retrieve", headers=AUTH,
                                 json={"user_id": str(uuid.uuid4()), "query": QUESTIONS[0][0], "document_ids": [str(d) for d in doc_ids]})).json()
        print(f"  isolation: another user querying these document ids gets {len(other['chunks'])} chunks -> {'PASS' if not other['chunks'] else 'FAIL'}")
        checks["user isolation"] = not other["chunks"]

        # ---------------------------------------------------------------- 4. end-to-end through the running Agent Core
        print("\n=== END-TO-END (Agent Core :8001 -> RAG Core -> Pinecone -> LLM :8003 -> FreeLLMAPI) ===")
        ok_ans, ok_neg = 0, 0
        for q, history, expect, cited in E2E + [(q, [], None, None) for q in UNANSWERABLE]:
            body = {"user_id": str(user), "conversation_id": str(uuid.uuid4()), "message_id": str(uuid.uuid4()), "message": q,
                    "document_ids": [str(d) for d in doc_ids], "history": history, "context": {}}
            answer, done, event = "", {}, None
            async with http.stream("POST", f"{AGENT}/api/v1/agent/stream", headers=AUTH, json=body) as resp:
                async for line in resp.aiter_lines():
                    if line.startswith("event:"):
                        event = line[6:].strip()
                    elif line.startswith("data:"):
                        data = json.loads(line[5:])
                        if event == "content.delta":
                            answer += data["delta"]
                        elif event in ("generation.completed", "error"):
                            done = {"event": event, **data}
            cites = [c["filename"] for c in done.get("citations", [])]
            md = done.get("metadata", {})
            if expect is not None:
                fact, cite = any(e in norm(answer) for e in expect), cited in cites
                ok_ans += fact and cite
                print(f"\nQ: {q}\n   fact={'PASS' if fact else 'FAIL'} citation={'PASS' if cite else 'FAIL'} cited={cites} "
                      f"strategy={md.get('retrieval_strategy')} chunks={md.get('retrieved_chunks')} rewritten={md.get('query_rewritten')} model={md.get('model')}")
            else:
                cleaned = re.sub(r"[*_`]", "", norm(answer))  # ignore Markdown emphasis
                declined = ABSTAIN_SENTENCE in cleaned or bool(REFUSAL.search(cleaned))
                ok_neg += declined
                print(f"\nQ (unanswerable): {q}\n   declined={'PASS' if declined else 'CHECK'} cited={cites} model={md.get('model')}")
            if done.get("event") == "error":
                print(f"   ERROR {done.get('code')}")
            print("   A:", " ".join(answer.split())[:240])
        print(f"\nE2E RESULT: answerable {ok_ans}/{len(E2E)} correct+cited | unanswerable {ok_neg}/{len(UNANSWERABLE)} declined")
        # LLM output varies run to run (FreeLLMAPI may route to different free models): allow one miss.
        checks["E2E answerable >= 6/7"] = ok_ans >= len(E2E) - 1
        checks["E2E unanswerable declined"] = ok_neg == len(UNANSWERABLE)
    finally:
        # ---------------------------------------------------------------- 5. cleanup
        print("\n=== CLEANUP ===")
        for did in docs:
            r = await http.delete(f"{RAG}/api/v1/rag/documents/{did}", headers=AUTH, params={"user_id": str(user)})
            if r.status_code not in (204, 404):
                print("  delete failed", did, r.status_code)
        deadline = time.time() + 120
        while True:
            count, total = await pinecone_stats(store, user)
            if count == 0 or time.time() > deadline:
                break
            await asyncio.sleep(3)
        async with rag_db.SessionLocal() as db:
            left = await db.scalar(select(func.count()).select_from(Chunk).where(Chunk.user_id == user))
        print(f"  benchmark namespace vectors={count} (index total={total}) rag_chunks left={left} -> {'CLEAN' if count == 0 and left == 0 else 'NOT CLEAN'}")
        checks["cleanup"] = count == 0 and left == 0
        await http.aclose()
        await store.aclose()
        await rag_db.engine.dispose()

    print("\n=== BASELINE GATE ===")
    for name, passed in checks.items():
        print(f"  {'PASS' if passed else 'FAIL'}  {name}")
    ok = len(checks) == 8 and all(checks.values())  # a missing check means a phase did not run
    print(f"  -> {'BASELINE MET' if ok else 'BASELINE NOT MET'}")
    return ok


if __name__ == "__main__":
    sys.exit(0 if asyncio.run(main()) else 1)
